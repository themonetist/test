#!/usr/bin/env python3
"""
build-leads-xlsx.py — turn a wassenger-backup.js output folder into an Excel
workbook with every lead and every conversation.

Usage:
    python3 build-leads-xlsx.py <backup-dir> <output.xlsx>

Sheets:
    Summary   — totals and open-lead count (formulas over the Leads sheet)
    Leads     — one row per phone number: name, labels, activity, last message,
                click-to-chat link and the FULL conversation text
    Messages  — one row per message: phone, name, date, direction, type, text,
                media file. Filter by number or keyword to see what was discussed.
"""

import json
import re
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

FONT = "Arial"
BALI = timezone(timedelta(hours=8))  # WITA
CELL_LIMIT = 32000  # Excel hard limit is 32,767 chars per cell

# Characters that are illegal in the XML inside an .xlsx. WhatsApp messages
# occasionally contain them (stray control codes, lone surrogates) and Excel
# refuses to open a file that has even one.
_ILLEGAL = re.compile(
    "[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f\ud800-\udfff￾￿]"
)


def clean(value):
    """Return a string Excel will accept: illegal XML chars removed, length capped."""
    if value is None:
        return None
    s = _ILLEGAL.sub("", str(value))
    if len(s) > CELL_LIMIT:
        s = s[: CELL_LIMIT - 40] + "\n[... truncated ...]"
    return s


def append_text_row(ws, values):
    """ws.append() but every string is cleaned and stored as text, never as a formula."""
    ws.append([clean(v) if isinstance(v, str) else v for v in values])
    for c in ws[ws.max_row]:
        if isinstance(c.value, str):
            c.data_type = "s"


def parse_ts(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return datetime.fromtimestamp(v / 1000 if v > 1e12 else v, tz=timezone.utc)
    try:
        return datetime.fromisoformat(str(v).replace("Z", "+00:00"))
    except ValueError:
        return None


def msg_ts(m):
    return parse_ts(m.get("date") or m.get("timestamp") or m.get("createdAt"))


def msg_body(m):
    return m.get("body") or m.get("message") or m.get("text") or (m.get("media") or {}).get("caption") or ""


def msg_from_me(m):
    if m.get("flow"):
        return m["flow"] == "outbound"
    return bool(m.get("fromMe") or m.get("outbound") or m.get("me"))


def naive_bali(dt):
    return dt.astimezone(BALI).replace(tzinfo=None) if dt else None


def style_header(ws, ncols):
    fill = PatternFill("solid", fgColor="1F3864")
    for col in range(1, ncols + 1):
        c = ws.cell(row=1, column=col)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"


def main(backup_dir, out_path):
    backup = Path(backup_dir)
    index = json.loads((backup / "index.json").read_text(encoding="utf-8"))
    entries = index["index"]

    wb = Workbook()

    # ------------------------------------------------------------------ Leads
    ws = wb.active
    ws.title = "Leads"
    lead_headers = [
        "Phone", "Name", "Labels", "Status", "Messages", "Inbound", "Outbound", "Media files",
        "First message (Bali)", "Last message (Bali)", "Last from", "Last message text",
        "WhatsApp link", "Conversation (full transcript)",
    ]
    ws.append(lead_headers)

    message_rows = []
    truncated = 0

    for e in entries:
        chat_file = backup / e["file"]
        data = json.loads(chat_file.read_text(encoding="utf-8")) if chat_file.exists() else {"messages": []}
        messages = sorted(data.get("messages", []), key=lambda m: (msg_ts(m) or datetime.min.replace(tzinfo=timezone.utc)))

        lines = []
        for m in messages:
            ts = msg_ts(m)
            when = naive_bali(ts).strftime("%Y-%m-%d %H:%M") if ts else "????-??-?? ??:??"
            who = "ME" if msg_from_me(m) else "THEM"
            mtype = m.get("type") or "text"
            media = m.get("media") or {}
            body = msg_body(m).replace("\r\n", "\n")
            tag = f"[{mtype}{' ' + media['filename'] if media.get('filename') else ''}] " if media or mtype not in ("text", "chat") else ""
            lines.append(f"{when} {who}: {tag}{body}")
            message_rows.append([
                f"+{e['phone']}", e.get("name"), naive_bali(ts), "Me" if msg_from_me(m) else "Them",
                mtype, body, media.get("filename") or (media.get("id") + "." + media.get("extension", "bin") if media.get("id") else None),
            ])

        transcript = "\n".join(lines)
        if len(transcript) > CELL_LIMIT:
            transcript = transcript[:CELL_LIMIT - 60] + "\n[... truncated, full text on the Messages sheet ...]"
            truncated += 1

        append_text_row(ws, [
            f"+{e['phone']}", e.get("name"), "; ".join(e.get("labels") or []), e.get("status"),
            e.get("messages", 0), e.get("inbound", 0), e.get("outbound", 0), e.get("media", 0),
            naive_bali(parse_ts(e.get("firstAt"))), naive_bali(parse_ts(e.get("lastAt"))),
            {"me": "Me", "them": "Them"}.get(e.get("lastFrom"), ""), e.get("lastBody"),
            f"https://wa.me/{e['phone']}", transcript,
        ])

    n_leads = ws.max_row - 1
    style_header(ws, len(lead_headers))
    widths = [16, 26, 24, 10, 10, 9, 9, 9, 18, 18, 9, 50, 30, 90]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            c.font = Font(name=FONT)
        row[8].number_format = "yyyy-mm-dd hh:mm"
        row[9].number_format = "yyyy-mm-dd hh:mm"
        row[11].alignment = Alignment(wrap_text=True, vertical="top")
        row[13].alignment = Alignment(wrap_text=True, vertical="top")
        link = row[12]
        link.hyperlink = link.value
        link.font = Font(name=FONT, color="0563C1", underline="single")
        ws.row_dimensions[row[0].row].height = 60
    if n_leads:
        tbl = Table(displayName="LeadsTable", ref=f"A1:{get_column_letter(len(lead_headers))}{ws.max_row}")
        tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tbl)

    # --------------------------------------------------------------- Messages
    wm = wb.create_sheet("Messages")
    msg_headers = ["Phone", "Name", "Date (Bali)", "From", "Type", "Message", "Media file"]
    wm.append(msg_headers)
    message_rows.sort(key=lambda r: (r[0], r[2] or datetime.min))
    for r in message_rows:
        append_text_row(wm, r)
    style_header(wm, len(msg_headers))
    for i, w in enumerate([16, 26, 18, 7, 10, 90, 36], start=1):
        wm.column_dimensions[get_column_letter(i)].width = w
    for row in wm.iter_rows(min_row=2, max_row=wm.max_row):
        for c in row:
            c.font = Font(name=FONT)
        row[2].number_format = "yyyy-mm-dd hh:mm"
        row[5].alignment = Alignment(wrap_text=True, vertical="top")
    if message_rows:
        tbl = Table(displayName="MessagesTable", ref=f"A1:{get_column_letter(len(msg_headers))}{wm.max_row}")
        tbl.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        wm.add_table(tbl)

    # ---------------------------------------------------------------- Summary
    wsum = wb.create_sheet("Summary", 0)
    last = ws.max_row
    rows = [
        ("Mandaya Energy — WhatsApp backup", None),
        ("Generated", index.get("generatedAt")),
        ("Wassenger device", index.get("device")),
        (None, None),
        ("Contacts (leads)", f"=COUNTA(Leads!A2:A{last})" if n_leads else 0),
        ("Contacts with at least one message", f"=COUNTIF(Leads!E2:E{last},\">0\")" if n_leads else 0),
        ("Total messages", f"=SUM(Leads!E2:E{last})" if n_leads else 0),
        ("Inbound messages", f"=SUM(Leads!F2:F{last})" if n_leads else 0),
        ("Outbound messages", f"=SUM(Leads!G2:G{last})" if n_leads else 0),
        ("Media files", f"=SUM(Leads!H2:H{last})" if n_leads else 0),
        ("Open: last message from them (needs a reply)", f"=COUNTIF(Leads!K2:K{last},\"Them\")" if n_leads else 0),
        ("Waiting: last message from us", f"=COUNTIF(Leads!K2:K{last},\"Me\")" if n_leads else 0),
        (None, None),
        ("How to use", "Leads: one row per number, full conversation in the last column. "
                       "Messages: one row per message, use the filter on Phone or search Message. "
                       "Click the WhatsApp link to open the chat on your phone or WhatsApp Web."),
        ("Note", f"{truncated} transcript(s) exceeded Excel's 32,767-character cell limit and were cut; "
                 "the Messages sheet always has the complete text." if truncated else
                 "All transcripts fit in their cells."),
        ("Source", "Wassenger REST API, exported with tools/wassenger-backup.js; times shown in Bali time (UTC+8)."),
    ]
    for label, value in rows:
        wsum.append([label, value])  # Summary values are formulas on purpose
    wsum["A1"].font = Font(name=FONT, bold=True, size=14)
    for row in wsum.iter_rows(min_row=2, max_row=wsum.max_row):
        row[0].font = Font(name=FONT, bold=True)
        row[1].font = Font(name=FONT)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    wsum.column_dimensions["A"].width = 44
    wsum.column_dimensions["B"].width = 90

    wb.save(out_path)
    print(f"wrote {out_path}: {n_leads} leads, {len(message_rows)} messages, {truncated} truncated transcripts")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
